"""
Helpers for building supplemental workbooks (Web/Burp and Nexpose) in memory for tests.

The sheet names and headers mirror the workbooks produced by Cyberwriter's data pipeline, so
the parsers are exercised against the same contract they see in production. Nothing here is a
test module; import it from the ``test_*`` files.
"""

# Standard Libraries
from io import BytesIO
from typing import Dict, Iterable, List, Optional, Sequence

# Django Imports
from django.core.files.uploadedfile import SimpleUploadedFile

# 3rd Party Libraries
import openpyxl
import xlsxwriter

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

NEXPOSE_UNIQUE_HEADERS = ["Risk", "Issue", "Impact", "Remediation", "Category"]
NEXPOSE_ALL_HEADERS = [
    "IP Address",
    "Hostname(s)",
    "Port",
    "Issue",
    "Impact",
    "Issue Details",
    "Evidence",
    "Remediation",
    "Risk",
    "Category",
]
WEB_UNIQUE_HEADERS = ["Issue", "Impact", "Background", "Fix", "Risk"]
WEB_ALL_HEADERS = [
    "Issue",
    "Impact",
    "Background",
    "Host",
    "Path",
    "Evidence",
    "Fix",
    "Detailed Remediation",
    "Risk",
]


def row(headers: Sequence[str], **values) -> List:
    """Build a row aligned to ``headers`` from keyword values keyed by a simplified header name."""

    def key(name: str) -> str:
        return "".join(ch for ch in name.lower() if ch.isalnum())

    lookup = {key(k): v for k, v in values.items()}
    return [lookup.get(key(header), "") for header in headers]


def build_workbook(sheets: Dict[str, Iterable[Sequence]]) -> bytes:
    """Write ``{sheet_name: rows}`` to an in-memory XLSX and return its bytes."""
    output = BytesIO()
    workbook = xlsxwriter.Workbook(output, {"in_memory": True})
    for name, rows in sheets.items():
        worksheet = workbook.add_worksheet(name)
        for row_index, values in enumerate(rows):
            for col_index, value in enumerate(values):
                if value is None or value == "":
                    continue
                if isinstance(value, bool):
                    worksheet.write_boolean(row_index, col_index, value)
                elif isinstance(value, (int, float)):
                    worksheet.write_number(row_index, col_index, value)
                else:
                    worksheet.write_string(row_index, col_index, str(value))
    workbook.close()
    return output.getvalue()


def default_nexpose_unique_rows() -> List[List]:
    """Three distinct issues; "Outdated OpenSSL" appears twice with mixed risk and an empty remediation."""
    h = NEXPOSE_UNIQUE_HEADERS
    return [
        row(
            h,
            risk="High",
            issue="Outdated OpenSSL",
            impact="Impact A",
            remediation="Upgrade OpenSSL",
            category="Crypto",
        ),
        row(
            h,
            risk="Medium",
            issue="SMB Signing Disabled",
            impact="Impact B",
            remediation="Enable SMB signing",
            category="SMB",
        ),
        row(
            h,
            risk="Low",
            issue="Outdated OpenSSL",
            impact="Impact A",
            remediation="",
            category="Crypto",
        ),
        row(
            h,
            risk="Low",
            issue="ICMP Timestamp",
            impact="Impact C",
            remediation="Block ICMP",
            category="Network",
        ),
    ]


def default_nexpose_all_rows() -> List[List]:
    """Hosts for the default unique issues; "ICMP Timestamp" intentionally has no hosts."""
    h = NEXPOSE_ALL_HEADERS
    return [
        row(
            h,
            ipaddress="10.0.0.1",
            hostnames="host-a",
            port=443,
            issue="Outdated OpenSSL",
            risk="High",
        ),
        row(
            h,
            ipaddress="10.0.0.2",
            hostnames="",
            port=443,
            issue="Outdated OpenSSL",
            risk="High",
        ),
        row(
            h,
            ipaddress="10.0.0.1",
            hostnames="host-a",
            port=8443,
            issue="Outdated OpenSSL",
            risk="High",
        ),
        row(
            h,
            ipaddress="",
            hostnames="host-c",
            port=445,
            issue="SMB Signing Disabled",
            risk="Medium",
        ),
    ]


def nexpose_workbook(
    unique_rows: Optional[List[List]] = None,
    all_rows: Optional[List[List]] = None,
    unique_headers: Sequence[str] = NEXPOSE_UNIQUE_HEADERS,
    all_headers: Sequence[str] = NEXPOSE_ALL_HEADERS,
    unique_sheet: str = "Unique Issues",
    all_sheet: str = "All Issues",
    extra_sheets: Optional[Dict[str, List[List]]] = None,
) -> bytes:
    sheets: Dict[str, List[List]] = {
        "Executive Summary": [["Total", 4]],
        unique_sheet: [list(unique_headers)]
        + (unique_rows if unique_rows is not None else default_nexpose_unique_rows()),
        all_sheet: [list(all_headers)]
        + (all_rows if all_rows is not None else default_nexpose_all_rows()),
    }
    if extra_sheets:
        sheets.update(extra_sheets)
    return build_workbook(sheets)


def default_web_unique_rows() -> List[List]:
    h = WEB_UNIQUE_HEADERS
    return [
        row(
            h,
            issue="Cross-site scripting (reflected)",
            impact="Impact",
            background="Bg",
            fix="Encode output",
            risk="High",
        ),
        row(
            h,
            issue="Cookie without HttpOnly flag set",
            impact="Impact",
            background="Bg",
            fix="Set HttpOnly",
            risk="Low",
        ),
    ]


def default_web_all_rows() -> List[List]:
    h = WEB_ALL_HEADERS
    return [
        row(
            h,
            issue="Cross-site scripting (reflected)",
            host="https://b.example",
            path="/search",
            fix="Encode output",
            risk="High",
        ),
        row(
            h,
            issue="Cross-site scripting (reflected)",
            host="https://a.example",
            path="/login",
            fix="Encode output",
            risk="High",
        ),
        row(
            h,
            issue="Cross-site scripting (reflected)",
            host="https://b.example",
            path="/search",
            fix="Encode output",
            risk="High",
        ),
        row(
            h,
            issue="Cookie without HttpOnly flag set",
            host="https://a.example",
            path="/",
            fix="Set HttpOnly",
            risk="Low",
        ),
    ]


def web_workbook(
    unique_rows: Optional[List[List]] = None,
    all_rows: Optional[List[List]] = None,
    unique_headers: Sequence[str] = WEB_UNIQUE_HEADERS,
    all_headers: Sequence[str] = WEB_ALL_HEADERS,
    unique_sheet: str = "Unique Issues",
    all_sheet: str = "All Issues",
    extra_sheets: Optional[Dict[str, List[List]]] = None,
) -> bytes:
    sheets: Dict[str, List[List]] = {
        "Executive Summary": [["Total", 4]],
        unique_sheet: [list(unique_headers)]
        + (unique_rows if unique_rows is not None else default_web_unique_rows()),
        all_sheet: [list(all_headers)]
        + (all_rows if all_rows is not None else default_web_all_rows()),
    }
    if extra_sheets:
        sheets.update(extra_sheets)
    return build_workbook(sheets)


def uploaded(
    name: str, data: bytes, content_type: str = XLSX_CONTENT_TYPE
) -> SimpleUploadedFile:
    return SimpleUploadedFile(name, data, content_type=content_type)


def read_workbook(data: bytes) -> openpyxl.Workbook:
    """Load generated XLSX bytes (e.g. a CAP export) for assertions."""
    return openpyxl.load_workbook(BytesIO(data))


def sheet_rows(worksheet) -> List[List]:
    """All rows of a worksheet as lists of cell values."""
    return [list(values) for values in worksheet.iter_rows(values_only=True)]
